import importlib.util
import csv
import json
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("build_facebook_workbook", ROOT / "scripts" / "build_facebook_workbook.py")
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class WorkbookTests(unittest.TestCase):
    def test_builds_eight_consistent_sheets(self):
        with tempfile.TemporaryDirectory() as value:
            out = Path(value)
            parent = {"record_id": "p", "comment_id": "p", "post_id": "post", "depth": 0,
                      "parent_comment_id": "post:post", "post_audit_status": "PASS", "body": "parent"}
            reply = {"record_id": "r", "comment_id": "r", "post_id": "post", "depth": 1,
                     "parent_comment_id": "p", "post_audit_status": "PASS", "body": "helpful reply text"}
            raw = [parent, reply]
            for name, rows in (("raw_comments.jsonl", raw), ("trusted_comments.jsonl", raw),
                               ("partial_candidates.jsonl", [])):
                (out / name).write_text("".join(json.dumps(row) + "\n" for row in rows), encoding="utf-8")
            with (out / "raw_comments.csv").open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=list(parent))
                writer.writeheader()
                writer.writerows(raw)
            (out / "conversation_map.json").write_text(json.dumps({"asins": ["B003ULL1NQ"], "layers": {"community_language": ["mobility"]}, "products": [], "sellersprite": {"status": "degraded"}}), encoding="utf-8")
            (out / "query_plan.json").write_text(json.dumps({"queries": [{"query": "dog mobility"}]}), encoding="utf-8")
            (out / "page_map.json").write_text(json.dumps([{"post_id": "post", "status": "PASS"}]), encoding="utf-8")
            evidence = out / "evidence" / "post"
            evidence.mkdir(parents=True)
            for name in ("audit.json", "comment_tree.json", "frontier.json", "action_ledger.json"):
                (evidence / name).write_text("{}", encoding="utf-8")
            (out / "manifest.json").write_text(json.dumps({"final_collected_count": 2,
                "ui_artifact_count": 0, "trusted_comment_count": 2, "reply_count": 1,
                "technical_duplicate_count": 0, "post_count": 1}), encoding="utf-8")
            result = MODULE.build(out)
            self.assertEqual(result["raw"], 2)
            self.assertTrue((out / "B003ULL1NQ_Facebook评论.xlsx").exists())
            import openpyxl
            workbook = openpyxl.load_workbook(out / "B003ULL1NQ_Facebook评论.xlsx", read_only=True)
            self.assertEqual(tuple(workbook.sheetnames), MODULE.SHEETS)
            self.assertEqual(workbook["Raw_Comments"].max_row - 1, 2)
            self.assertEqual(workbook["Trusted_Comments"].max_row - 1, 2)
            self.assertEqual(workbook["Partial_Candidates"].max_row - 1, 0)
            summary = {row[0]: row[1] for row in workbook["Run_Summary"].iter_rows(min_row=2, values_only=True)}
            for metric in ("ui_artifact_count", "trusted_comment_count", "reply_count", "technical_duplicate_count"):
                self.assertIn(metric, summary)
            quality = {row[0]: row[1] for row in workbook["Quality_Gate"].iter_rows(min_row=2, values_only=True)}
            workbook.close()
            self.assertTrue(all(value == "PASS" for value in quality.values()))

    def test_quality_gate_rejects_ui_artifact_in_trusted(self):
        raw = [{"record_id": "ui", "comment_id": "ui", "post_id": "p", "depth": 0,
                "body_clean": "38周", "ui_artifact": True, "record_kind": "ui_artifact"}]
        checks = {row["check"]: row["status"] for row in MODULE.quality_rows(
            raw, raw, [], [{"post_id": "p"}], {"final_collected_count": 1,
            "ui_artifact_count": 1, "trusted_comment_count": 1, "post_count": 1})}
        self.assertEqual(checks["trusted_has_no_ui_artifacts"], "FAIL")
        self.assertEqual(checks["trusted_has_no_ui_only_text"], "FAIL")


if __name__ == "__main__":
    unittest.main()
